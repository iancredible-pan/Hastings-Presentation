"""Build matching PPTX/PDF slides and a printable English speaker-notes booklet.

Usage: python source/build_presentation.py
Dependencies: python-pptx, reportlab, matplotlib, Pillow. No network required.
"""
from pathlib import Path
import os
import re
import json
import hashlib
import math

ROOT = Path(__file__).resolve().parents[1]
CACHE = ROOT / '.build'
CACHE.mkdir(exist_ok=True)
os.environ.setdefault('MPLCONFIGDIR', str(CACHE / 'matplotlib'))

import matplotlib
matplotlib.use('Agg')
from matplotlib.figure import Figure
from matplotlib.backends.backend_agg import FigureCanvasAgg
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import ImageReader
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from xml.sax.saxutils import escape

FONTDIR = Path('/usr/share/fonts/truetype/liberation2')
if not FONTDIR.exists():
    FONTDIR = Path('/usr/share/fonts/truetype/liberation')
for name, file in [('Body','LiberationSans-Regular.ttf'),('Bold','LiberationSans-Bold.ttf'),
                   ('Italic','LiberationSans-Italic.ttf'),('Mono','LiberationMono-Regular.ttf')]:
    pdfmetrics.registerFont(TTFont(name, str(FONTDIR / file)))
pdfmetrics.registerFontFamily('Body', normal='Body', bold='Bold', italic='Italic', boldItalic='Bold')

W, H = 13.333333, 7.5
NAVY = '14283F'
TEAL = '007F78'
GOLD = 'EABF62'
INK = '192D43'
MUTED = '586B7C'
BG = 'F6F7F3'
PALE = 'E7F1EC'
WHITE = 'FFFFFF'
LINE = 'D4DED9'
RED = 'AE4939'


def wrap(text, width, size, font):
    """Use the same measured lines in the editable deck and PDF."""
    out = []
    for line in text.split('\n'):
        if not line:
            out.append('')
            continue
        if pdfmetrics.stringWidth(line, font, size) <= width * 72:
            out.append(line)
            continue
        words, current = line.split(), ''
        for word in words:
            candidate = (current + ' ' + word).strip()
            if current and pdfmetrics.stringWidth(candidate, font, size) > width * 72:
                out.append(current)
                current = word
            else:
                current = candidate
        out.append(current)
    return out


def math_image(tex, size=30, color=INK):
    key = hashlib.sha256(f'{tex}/{size}/{color}'.encode()).hexdigest()[:16]
    path = CACHE / f'eq-{key}.png'
    if not path.exists():
        matplotlib.rcParams['mathtext.fontset'] = 'dejavusans'
        fig = Figure(figsize=(15, 2), dpi=220)
        FigureCanvasAgg(fig)
        artist = fig.text(0, 0, '$' + tex + '$', size=size, color='#' + color)
        fig.canvas.draw()
        bbox = artist.get_window_extent(fig.canvas.get_renderer()).expanded(1.025, 1.15)
        fig.savefig(path, transparent=True, dpi=220,
                    bbox_inches=bbox.transformed(fig.dpi_scale_trans.inverted()), pad_inches=.015)
    with Image.open(path) as img:
        return path, img.width / 220, img.height / 220


class Slide:
    def __init__(self, title, section, presenter='', minutes=0, source='', dark=False, backup=False):
        self.title, self.section = title, section
        self.presenter, self.minutes, self.source = presenter, minutes, source
        self.dark, self.backup = dark, backup
        self.ops = []
        self.bg = NAVY if dark else BG
        self.fg = WHITE if dark else INK
        self.rect(.45, .46, .12, .16, TEAL if not dark else GOLD)
        self.text(section.upper(), .7, .42, 11.6, size=11, color=GOLD if dark else TEAL, bold=True)
        self.text(title, .7, .91, 12, size=31, color=self.fg, bold=True)
        self.line(.7, 6.98, 12.62, 6.98, color='385066' if dark else LINE, width=.7)

    def rect(self, x, y, w, h, fill, radius=False, stroke=None):
        self.ops.append(dict(kind='rect',x=x,y=y,w=w,h=h,fill=fill,radius=radius,stroke=stroke))

    def circle(self, x, y, diameter, fill, stroke=None):
        self.ops.append(dict(kind='circle',x=x,y=y,w=diameter,h=diameter,fill=fill,stroke=stroke))

    def text(self, text, x, y, w, size=22, color=None, bold=False, align='left', mono=False, h=None):
        font = 'Mono' if mono else ('Bold' if bold else 'Body')
        lines = wrap(text, w, size, font)
        height = max(.2, len(lines) * size * 1.22 / 72 + .06)
        if h is not None and height > h + .035:
            raise ValueError(f'Text overflow ({height:.2f}>{h}): {text}')
        self.ops.append(dict(kind='text',text=text,lines=lines,x=x,y=y,w=w,h=height,
                             size=size,color=color or self.fg,bold=bold,align=align,mono=mono))
        return height

    def eq(self, tex, x, y, w=11.8, size=30, color=None, max_h=1.5, center=False):
        path, ew, eh = math_image(tex, size, color or self.fg)
        scale = min(1, w / ew, max_h / eh)
        ew, eh = ew * scale, eh * scale
        self.ops.append(dict(kind='image',path=str(path),x=x+(w-ew)/2 if center else x,
                             y=y,w=ew,h=eh,tex=tex))
        return eh

    def line(self, x1,y1,x2,y2,color=TEAL,width=2,arrow=False):
        self.ops.append(dict(kind='line',x1=x1,y1=y1,x2=x2,y2=y2,color=color,width=width,arrow=arrow))

    def label(self, text, x, y, w=3, color=TEAL):
        self.text(text.upper(),x,y,w,size=11,color=color,bold=True)

    def card(self, x,y,w,h, fill=WHITE):
        self.rect(x,y,w,h,fill,radius=True)

    def callout(self, text, y=6.04, fill=PALE, color=TEAL, size=20):
        self.rect(.7,y,11.92,.63,fill,radius=True)
        self.text(text,.91,y+.14,11.5,size=size,color=color,bold=True,h=.48)

    def matrix(self, label, entries, x, y, cellw=.82, cellh=.66, size=28, color=None):
        color = color or self.fg
        self.eq(label+'=',x,y+.23,w=1.1,size=size)
        left, top = x+1.25,y
        width,height=2*cellw,2*cellh
        for bx,side in [(left,-1),(left+width,1)]:
            self.line(bx,top,bx,top+height,color=color,width=1.5)
            self.line(bx,top,bx-side*.14,top,color=color,width=1.5)
            self.line(bx,top+height,bx-side*.14,top+height,color=color,width=1.5)
        for row in range(2):
            for col in range(2):
                self.eq(entries[row][col],left+col*cellw+.07,top+row*cellh+.04,
                        w=cellw-.14,size=size,color=color,max_h=cellh-.06,center=True)

    def graph(self, top='1/3', bottom='1', selfloop=False, flows=False, y=3.25):
        x0,x1=3.1,9.45
        diameter=1.25
        self.line(x0+diameter+.1,y+.29,x1-.1,y+.29,arrow=True)
        self.line(x1-.1,y+.96,x0+diameter+.1,y+.96,arrow=True)
        self.circle(x0,y,diameter,TEAL)
        self.circle(x1,y,diameter,NAVY)
        self.text('0',x0,y+.31,diameter,size=33,color=WHITE,bold=True,align='center')
        self.text('1',x1,y+.31,diameter,size=33,color=WHITE,bold=True,align='center')
        self.eq(top,5.3,y-.39,w=2.8,size=28,center=True)
        self.eq(bottom,5.3,y+1.02,w=2.8,size=28,center=True)
        if selfloop:
            self.line(x0+.2,y-.04,x0+.2,y-.7,color=TEAL)
            self.line(x0+.2,y-.7,x0-.8,y-.7,color=TEAL)
            self.line(x0-.8,y-.7,x0-.8,y+.65,color=TEAL)
            self.line(x0-.8,y+.65,x0-.04,y+.65,color=TEAL,arrow=True)
            self.eq(r'2/3',x0-1.5,y-.35,w=.8,size=23)
        if flows:
            self.text('target mass 3/4',x0-.75,y+1.6,2.8,size=18,color=TEAL,align='center')
            self.text('target mass 1/4',x1-.75,y+1.6,2.8,size=18,color=MUTED,align='center')


slides=[]
def add(*a,**kw):
    s=Slide(*a,**kw);slides.append(s);return s

s=add('Sampling by designing a Markov chain','Hastings · 1970','A',1,'Hastings (1970), §§1–2.2',dark=True)
s.text('Monte Carlo sampling methods using Markov\nchains and their applications',.73,1.94,10.9,size=25,color='DEE8EA')
s.rect(.72,3.16,11.88,1.42,'203B52',radius=True)
s.text('How can we build a chain that preserves\na distribution of our choice?',1.02,3.45,11.2,size=29,color=WHITE,bold=True)
s.text('W. K. Hastings · Biometrika 57(1), 97–109',.74,5.13,11.7,size=17,color='CCD8DF')
s.text('Assigned reading: Sections 1, 2.1 and 2.2',.74,5.58,11.7,size=17,color='CCD8DF')
s.text('Presenter A  +  Presenter B',.74,6.19,11.7,size=17,color=GOLD,bold=True)

s=add('We want expectations. We need samples.','01 · The problem','A',2.5,'§1 · pp. 97–98')
s.eq(r'I=\mathbb{E}_{p}[f(X)]=\int f(x)\,p(x)\,dx',.9,1.9,w=11.5,size=34)
for x,title,body in [(0.7,'The quantity','A mean, probability, or other summary of interest.'),
                      (4.76,'The obstacle','High-dimensional integrals can be hard to evaluate.'),
                      (8.82,'The opportunity','Density values may be easier to compute than samples.')]:
    s.card(x,3.25,3.8,2.14)
    s.text(title,x+.23,3.53,3.34,size=23,bold=True,color=TEAL)
    s.text(body,x+.23,4.13,3.34,size=21)
s.callout('Hastings changes how we obtain the samples.')

s=add('Keep the average. Change the sampling mechanism.','01 · The problem','A',2.5,'§§1–2.1 · pp. 97–98')
s.card(.7,1.9,5.8,3.2)
s.card(6.76,1.9,5.85,3.2,PALE)
s.label('Ordinary Monte Carlo',.97,2.16,5)
s.eq(r'X_1,\ldots,X_N\ \overset{\mathrm{iid}}{\sim}\ p',1,2.77,w=5,size=26)
s.text('Draw independently from the target.',1,3.74,4.95,size=22)
s.label('Markov chain Monte Carlo',7.02,2.16,5)
s.eq(r'X_0\rightarrow X_1\rightarrow\cdots\rightarrow X_N',7.04,2.77,w=5.2,size=26)
s.text('Design transitions that preserve the target.',7.04,3.74,4.95,size=22)
s.eq(r'\widehat I_N=\frac{1}{N}\sum_{t=1}^{N} f(X_t)',.94,5.34,w=11.3,size=27,max_h=.89,center=True)
s.text('Dependent samples require care when assessing Monte Carlo error.',.94,6.48,11.3,size=18,color=MUTED,align='center')

s=add('One target will guide the entire construction.','02 · The target','A',2,'Teaching example · not an original example from the paper')
s.label('Unnormalized weights',.9,1.93,5)
s.eq(r'w_0=3\qquad w_1=1',.9,2.43,w=5.4,size=34)
s.label('Desired stationary probabilities',7.07,1.93,5)
s.eq(r'\pi=(3/4,\ 1/4)',7.07,2.43,w=5.4,size=34)
s.rect(1.04,3.67,8.4,1.02,TEAL,radius=True)
s.rect(9.44,3.67,2.8,1.02,NAVY,radius=True)
s.text('STATE 0 · 75%',1.3,3.96,7.85,size=25,color=WHITE,bold=True,align='center')
s.text('STATE 1 · 25%',9.57,4.01,2.55,size=18,color=WHITE,bold=True,align='center')
s.text('Desired long-run occupation: three times as much time in state 0.',1.02,5.18,11.25,size=23)
s.callout('Which transitions will produce this target?')

s=add('A stationary distribution survives one transition.','02 · The target','A',3,'§2.1 · p. 98')
s.eq(r'p_{ij}=\Pr(X_{t+1}=j\mid X_t=i)',.95,1.9,w=11.3,size=30)
s.eq(r'p_{ij}\geq0\qquad\sum_jp_{ij}=1',.95,2.7,w=11.3,size=27)
for x,label in [(1.15,'Current distribution'),(8.54,'Next distribution')]:
    s.card(x,3.8,3.62,1.46,PALE)
    s.text(label,x+.15,4.01,3.32,size=18,color=TEAL,align='center')
    s.eq(r'\pi' if x<2 else r'\pi P',x+.15,4.41,w=3.32,size=31,center=True)
s.line(5.04,4.53,8.2,4.53,arrow=True)
s.text('one step using P',5.11,3.94,2.85,size=18,align='center')
s.rect(4.5,5.51,4.3,.68,PALE,radius=True)
s.eq(r'\pi P=\pi',4.48,5.57,w=4.3,size=34,center=True)
s.text('The distribution stays the same. Individual states can still change.',.95,6.39,11.5,size=19,color=MUTED,align='center')

s=add('Balance probability flow in both directions.','02 · The target','A',2.5,'§2.2 · p. 99 · Eq. (4)')
s.eq(r'\pi_i p_{ij}=\pi_j p_{ji}',3.68,1.87,w=6,size=35,center=True)
s.graph(top=r'\pi_0p_{01}',bottom=r'\pi_1p_{10}',flows=True,y=3.12)
s.text('Detailed balance is a sufficient condition for stationarity.',.93,5.91,11.3,size=23,bold=True,color=TEAL,align='center')
s.text('Equal weighted flows do not require equal transition probabilities.',.93,6.44,11.3,size=18,color=MUTED,align='center')

s=add('A convenient proposal may have the wrong target.','03 · The construction','A',2.5,'§2.2 · proposal mechanism; teaching example')
s.matrix('Q',[[r'0',r'1'],[r'1',r'0']],.88,1.95)
s.text('Always propose the other state.',5.0,2.22,6.95,size=25,bold=True,color=TEAL)
s.text('Easy to simulate. Not yet the chain we want.',5.0,2.89,6.95,size=22)
s.graph(top=r'q_{01}=1',bottom=r'q_{10}=1',y=4.08)
s.callout('Q preserves (1/2, 1/2). Our target is (3/4, 1/4).',y=6.12,size=20)

s=add('A proposal becomes a move only if we accept it.','03 · The construction','A',3,'§2.2 · pp. 99–100 · Eq. (5)')
for x,n,title,body in [(0.7,'1','PROPOSE','Draw a candidate j from row i of Q.'),
                          (4.76,'2','ACCEPT OR REJECT','Accept with probability α. Otherwise remain at i.'),
                          (8.82,'3','RECORD','Record the next state, including repeated states.')]:
    s.card(x,1.94,3.79,2.45)
    s.text(n,x+.22,2.1,.6,size=32,bold=True,color=TEAL)
    s.label(title,x+.23,2.78,3.3)
    s.text(body,x+.23,3.16,3.3,size=21)
s.eq(r'p_{ij}=q_{ij}\alpha_{ij}\quad(i\ne j)',1.02,4.76,w=11.2,size=32)
s.eq(r'p_{ii}=1-\sum_{j\ne i}q_{ij}\alpha_{ij}',1.02,5.65,w=11.2,size=29)
s.text('Staying can include a rejected move or a proposal of the current state.',1.04,6.48,11.1,size=17,color=MUTED)

s=add('We have the proposal. Now design the correction.','Presenter handoff','A',2,'Construction checklist',dark=True)
for y,n,t,sub in [(1.96,'01','Choose the target','Done: weights 3 and 1.'),
                  (3.04,'02','Choose a proposal','Done: propose switching states.'),
                  (4.12,'03','Choose acceptance probabilities','Next: preserve the desired probability flow.')]:
    s.text(n,.95,y,.8,size=29,bold=True,color=GOLD)
    s.text(t,2.01,y,10.15,size=26,bold=True)
    s.text(sub,2.01,y+.48,10.15,size=20,color='CCDAE1')
s.rect(.78,5.61,11.78,.79,'203B52',radius=True)
s.text('How should acceptance compensate for the proposal?',1.02,5.84,11.3,size=24,color=GOLD,bold=True)

s=add('Accept according to target mass and proposal flow.','03 · The construction','B',3,'§2.2 · common Metropolis–Hastings choice')
s.eq(r'r_{ij}=\frac{\pi_jq_{ji}}{\pi_iq_{ij}}',1.02,1.96,w=4.4,size=35)
s.eq(r'\alpha_{ij}=\min(1,r_{ij})',6.13,2.07,w=6.2,size=35)
s.card(.73,3.4,5.78,1.92)
s.card(6.76,3.4,5.84,1.92,PALE)
s.text('r ≥ 1',1.0,3.65,5.2,size=27,bold=True,color=TEAL)
s.text('Accept the proposed move.',1.0,4.31,5.1,size=23)
s.text('r < 1',7.02,3.65,5.2,size=27,bold=True,color=TEAL)
s.text('Accept with probability r.',7.02,4.31,5.1,size=23)
s.eq(r'q_{ij}=q_{ji}\quad\Longrightarrow\quad\alpha_{ij}=\min(1,\pi_j/\pi_i)',1.02,5.73,w=11.1,size=26)
s.text('The reverse/forward proposal ratio corrects an asymmetric proposal.',1.02,6.47,11.2,size=18,color=MUTED)

s=add('A short correctness check: the same flow both ways.','03 · The construction','B',2.5,'§2.2 · detailed balance verification')
s.label('Forward flow',.96,1.96,5)
s.eq(r'\pi_iq_{ij}\alpha_{ij}=\min(\pi_iq_{ij},\pi_jq_{ji})',.96,2.4,w=11.45,size=32)
s.label('Reverse flow',.96,3.49,5)
s.eq(r'\pi_jq_{ji}\alpha_{ji}=\min(\pi_jq_{ji},\pi_iq_{ij})',.96,3.93,w=11.45,size=32)
s.rect(.79,5.04,11.74,1.23,PALE,radius=True)
s.eq(r'\mathrm{Detailed\ balance}\quad\Longrightarrow\quad\pi P=\pi',1.04,5.36,w=11.1,size=29,center=True)
s.text('A sufficient construction; stationarity does not require reversibility.',.98,6.48,11.25,size=17,color=MUTED)

s=add('Hastings gives a family of acceptance rules.','03 · The construction','B',2,'§2.2 · pp. 99–100 · Eq. (6)')
s.eq(r'\alpha_{ij}=\frac{s_{ij}}{1+\pi_iq_{ij}/(\pi_jq_{ji})}',1.04,1.91,w=11.2,size=34,center=True)
s.text('Choose symmetric s, with acceptance probabilities between 0 and 1.',1.05,3.01,11.2,size=22,align='center')
for x,title,rule in [(.77,'Barker',r'\alpha_{ij}=\frac{r_{ij}}{1+r_{ij}}'),
                     (6.76,'Metropolis–Hastings',r'\alpha_{ij}=\min(1,r_{ij})')]:
    s.card(x,3.77,5.81,1.83)
    s.text(title,x+.26,4.02,5.29,size=24,bold=True,color=TEAL)
    s.eq(rule,x+.26,4.66,w=5.26,size=28)
s.callout('Detailed balance does not uniquely determine the acceptance rule.',y=6.07,size=19)

s=add('The normalizing constant disappears from the ratio.','04 · The practical advantage','B',2.5,'§1 · p. 98; §2.2 · test ratio')
s.eq(r'\pi_i=\frac{w_i}{Z}\qquad Z=\sum_k w_k',1.04,1.93,w=11.2,size=32)
s.eq(r'r_{ij}=\frac{(w_j/Z)q_{ji}}{(w_i/Z)q_{ij}}=\frac{w_jq_{ji}}{w_iq_{ij}}',1.04,3.1,w=11.2,size=35)
s.card(.79,4.61,11.74,1.15,PALE)
s.text('Evaluate weights and proposal probabilities.\nNo need to compute the potentially difficult sum or integral Z.',1.04,4.82,11.2,size=23,color=TEAL,bold=True)
s.text('This removes normalization work; it does not guarantee fast exploration.',1.04,6.39,11.2,size=19,color=MUTED)

s=add('For our target, one direction needs more rejection.','05 · The worked example','B',2.5,'Teaching example · weights (3, 1), symmetric swap proposal')
s.text('Weights: 3 and 1.   Proposal: always switch.',1.02,1.93,11.2,size=23,color=MUTED)
for x,title,tex,body in [(.76,'From 0 to 1',r'\alpha_{01}=\min(1,1/3)=1/3','Accept one third of proposals.\nOtherwise stay in state 0.'),
                       (6.76,'From 1 to 0',r'\alpha_{10}=\min(1,3)=1','Accept every proposal.\nLeave state 1 immediately.')]:
    s.card(x,2.65,5.81,2.92)
    s.text(title,x+.27,2.94,5.25,size=24,bold=True,color=TEAL)
    s.eq(tex,x+.27,3.62,w=5.25,size=27)
    s.text(body,x+.27,4.53,5.25,size=22)
s.callout('State 0 gains its larger occupation through repeated visits and stays.',y=6.06,size=19)

s=add('Construct P, then verify the distribution we wanted.','05 · The worked example','B',3,'Teaching example · all entries and flows computed exactly')
s.matrix('P',[[r'\frac{2}{3}',r'\frac{1}{3}'],[r'1',r'0']],.95,1.94,cellw=1.0,cellh=.76,size=30)
s.eq(r'(3/4,\ 1/4)P=(3/4,\ 1/4)',5.31,2.3,w=6.5,size=27)
s.label('Check weighted flows',5.35,3.16,6)
s.eq(r'\frac{3}{4}\cdot\frac{1}{3}=\frac{1}{4}\cdot1=\frac{1}{4}',5.35,3.62,w=6.4,size=29)
s.graph(top=r'1/3',bottom=r'1',selfloop=True,y=4.97)

s=add('Rejected proposals are part of the sample.','05 · The worked example','B',2.5,'Teaching example · audience question')
s.text('What happens if we keep only accepted moves?',.95,1.91,11.4,size=27,bold=True,color=TEAL)
s.label('Illustrative recorded trajectory · includes every time step',.96,2.77,11.5)
states=[0,0,0,1,0,0,1,0,0,0,1,0]
for i,state in enumerate(states):
    x=1.02+i*.94
    s.circle(x,3.37,.64,TEAL if state==0 else NAVY)
    s.text(str(state),x,3.49,.64,size=22,color=WHITE,bold=True,align='center')
s.label('If rejected repeats are deleted',.96,4.45,11.5,color=RED)
s.eq(r'0\rightarrow1\rightarrow0\rightarrow1\rightarrow\cdots',1.04,4.91,w=8.3,size=30)
s.text('1 : 1',10.0,4.97,2.25,size=32,color=RED,bold=True)
s.callout('In this example, accepted moves alternate: the target 3 : 1 is lost.',y=6.05,size=19)

s=add('A small algorithm, with conditions we must check.','06 · Put it together','B',1.5,'§§2.1–2.2 · finite-state setting')
s.card(.75,1.91,7.51,4.71,WHITE)
code='choose X in the target support\nrepeat N times:\n    draw candidate J from Q[X, :]\n    draw U uniformly in (0, 1)\n    if U < alpha(X, J):\n        X = J\n    record X'
s.text(code,1.03,2.18,6.95,size=21,mono=True)
s.text('Use the MH rule above; set α(X, X) = 1.',1.03,5.9,6.95,size=18,color=TEAL)
s.label('Our finite-state example',8.63,2.05,3.8)
s.text('Both states communicate.\n\nState 0 has a self-loop.\n\nThe chain is irreducible and aperiodic.',8.63,2.61,3.67,size=22)

s=add('Design the transitions. Preserve the target.','06 · Takeaways','B',1.5,'Hastings (1970), §§1–2.2',dark=True)
for i,(title,body) in enumerate([
    ('Construct','Choose a proposal and an acceptance rule.'),
    ('Preserve','Balance weighted flows to obtain πP = π.'),
    ('Compute','Use unnormalized weights; Z cancels.'),
    ('Record','Keep every state, including rejected moves.')]):
    y=1.95+i*1.02
    s.text(f'0{i+1}',.94,y,.77,size=26,bold=True,color=GOLD)
    s.text(title,1.98,y,2.3,size=25,bold=True)
    s.text(body,4.3,y+.03,8.0,size=22,color='DCE6EA')
s.text('Which part of the construction makes the target distribution stationary?',.98,6.28,11.4,size=20,color=GOLD,bold=True)

s=add('An asymmetric proposal needs the correction factor.','Backup · Asymmetric proposal',source='§2.2 · teaching variation',backup=True)
s.matrix('Q',[[r'\frac{1}{2}',r'\frac{1}{2}'],[r'1',r'0']],.96,1.95,cellw=.94,cellh=.77,size=29)
s.eq(r'\alpha_{01}=\min\!\left(1,\frac{1\cdot1}{3\cdot(1/2)}\right)=\frac{2}{3}',5.2,2.0,w=7.0,size=27)
s.eq(r'\alpha_{10}=\min\!\left(1,\frac{3\cdot(1/2)}{1\cdot1}\right)=1',5.2,3.16,w=7.0,size=27)
s.eq(r'p_{01}=\frac{1}{2}\cdot\frac{2}{3}=\frac{1}{3}\qquad p_{10}=1',1.0,4.6,w=11.3,size=29)
s.callout('The corrected proposal produces the same P as our main example.',y=6.04,size=19)

s=add('What can the symmetric factor s actually be?','Backup · Hastings family',source='§2.2 · Eq. (6); assume positive pairwise flows',backup=True)
s.eq(r'A=\pi_iq_{ij}>0\qquad B=\pi_jq_{ji}>0\qquad r=B/A',.97,1.9,w=11.4,size=27)
s.eq(r'\alpha_{ij}=s\frac{B}{A+B}\qquad\alpha_{ji}=s\frac{A}{A+B}',.97,2.87,w=11.4,size=30)
s.eq(r'0<s\leq\frac{A+B}{\max(A,B)}=1+\min(r,1/r)',.97,3.89,w=11.4,size=29)
s.card(.79,5.06,11.74,1.41,PALE)
s.eq(r's=1\quad\Rightarrow\quad\mathrm{Barker}:\ \alpha=\frac{r}{1+r}',1.04,5.21,w=11.1,size=24)
s.eq(r's=1+\min(r,1/r)\quad\Rightarrow\quad\mathrm{MH}:\ \alpha=\min(1,r)',1.04,5.84,w=11.1,size=24)

s=add('Stationarity, convergence and independence differ.','Backup · Three distinct ideas',source='§2.1; finite-state Markov chain background',backup=True)
for i,(title,eq,body) in enumerate([
    ('Stationarity',r'\pi P=\pi','Starting in π preserves π at every time.'),
    ('Marginal convergence',r'\mu P^t\longrightarrow\pi','For finite chains: irreducibility plus aperiodicity suffices.'),
    ('Independence',r'\Pr(X_{t+1}=j\mid X_t=i)=\Pr(X_{t+1}=j)','Stationarity alone does not remove dependence.')]):
    y=1.92+i*1.43
    s.card(.77,y,11.8,1.24)
    s.text(title,1.0,y+.18,3.4,size=22,bold=True,color=TEAL)
    s.eq(eq,4.55,y+.12,w=7.55,size=23,max_h=.45)
    s.text(body,4.55,y+.74,7.55,size=17)
s.text('Finite irreducibility is enough for long-run time averages; aperiodicity is not required for that result.',.98,6.44,11.3,size=17,color=MUTED)

s=add('Dependence changes the uncertainty of an average.','Backup · Monte Carlo error',source='§2.1 · pp. 98–99; stationarity and finite variance assumed',backup=True)
s.eq(r'\mathrm{Var}(\bar Y_N)=\frac{\sigma_Y^2}{N}\left[1+2\sum_{k=1}^{N-1}\left(1-\frac{k}{N}\right)\rho_k\right]',.97,1.95,w=11.4,size=28)
s.eq(r'Y_t=f(X_t)\qquad\rho_k=\mathrm{Corr}(Y_t,Y_{t+k})',.97,3.29,w=11.4,size=25)
s.card(.78,4.45,11.78,1.54,PALE)
s.text('Positive correlations can increase uncertainty.\nNegative correlations can reduce it.',1.05,4.72,11.2,size=25,bold=True,color=TEAL)
s.text('In our example, f(x) = x has lag-k correlation (−1/3)^k under stationarity.',1.03,6.37,11.2,size=19,color=MUTED)

s=add('Sources and scope','Backup · References',source='Primary source and course guide',backup=True)
s.text('Primary source',1.0,1.98,11.2,size=24,bold=True,color=TEAL)
s.text('Hastings, W. K. (1970). Monte Carlo sampling methods using Markov chains and their applications. Biometrika, 57(1), 97–109.',1.0,2.5,11.2,size=23)
s.text('https://doi.org/10.1093/biomet/57.1.97',1.0,3.65,11.2,size=19,color=MUTED)
s.text('Assigned scope',1.0,4.31,11.2,size=24,bold=True,color=TEAL)
s.text('Sections 1, 2.1 and 2.2 (printed pp. 97–101).\nThe two-state construction is a teaching example, not an original example in the paper.',1.0,4.87,11.2,size=22)
s.text('Course guide: construct a chain with prescribed stationarity; explain acceptance, normalization cancellation and one simple example.',1.0,6.16,11.2,size=17,color=MUTED)


def notes_blocks():
    path=ROOT/'speaker_notes.md'
    if not path.exists():
        return {}
    text=path.read_text()
    matches=list(re.finditer(r'^## Slide (\d+)\b[^\n]*',text,re.M))
    return {int(m.group(1)):text[m.end(): matches[i+1].start() if i+1<len(matches) else len(text)].strip()
            for i,m in enumerate(matches)}


def save_slides():
    prs=Presentation()
    prs.slide_width=Inches(W);prs.slide_height=Inches(H)
    prs.core_properties.title='Hastings (1970): Sampling by designing a Markov chain'
    prs.core_properties.subject='Two-presenter, 42-minute presentation; assigned Sections 1, 2.1 and 2.2'
    prs.core_properties.author='Presenter A and Presenter B'
    prs.core_properties.keywords='Hastings; MCMC; Metropolis-Hastings; stationary distribution'
    pdf=canvas.Canvas(str(ROOT/'Hastings_1970_slides.pdf'),pagesize=(W*72,H*72))
    pdf.setTitle(prs.core_properties.title)
    notes=notes_blocks()
    for idx,s in enumerate(slides,1):
        ps=prs.slides.add_slide(prs.slide_layouts[6])
        ps.background.fill.solid();ps.background.fill.fore_color.rgb=RGBColor.from_string(s.bg)
        pdf.setFillColor(HexColor('#'+s.bg));pdf.rect(0,0,W*72,H*72,fill=1,stroke=0)
        footer_color='CAD7DE' if s.dark else MUTED
        s.text(s.source,.73,7.12,10.4,size=9,color=footer_color)
        owner=(f'{s.presenter}  ·  ' if s.presenter else 'BACKUP  ·  ')
        s.text(owner+f'{idx:02d}'+('/18' if idx<=18 else ''),11.12,7.1,1.5,size=10,color=footer_color,align='right')
        for op in s.ops:
            kind=op['kind']
            if kind in ('rect','circle'):
                shape=ps.shapes.add_shape(MSO_SHAPE.OVAL if kind=='circle' else (MSO_SHAPE.ROUNDED_RECTANGLE if op['radius'] else MSO_SHAPE.RECTANGLE),
                    Inches(op['x']),Inches(op['y']),Inches(op['w']),Inches(op['h']))
                shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(op['fill'])
                if op.get('stroke'):
                    shape.line.color.rgb=RGBColor.from_string(op['stroke'])
                else: shape.line.fill.background()
                if kind=='rect' and op['radius']:
                    shape.adjustments[0]=.08
                pdf.setFillColor(HexColor('#'+op['fill']))
                x,y,w,h=op['x']*72,(H-op['y']-op['h'])*72,op['w']*72,op['h']*72
                if kind=='circle': pdf.ellipse(x,y,x+w,y+h,fill=1,stroke=0)
                elif op['radius']: pdf.roundRect(x,y,w,h,7,fill=1,stroke=0)
                else: pdf.rect(x,y,w,h,fill=1,stroke=0)
            elif kind=='text':
                tb=ps.shapes.add_textbox(Inches(op['x']),Inches(op['y']),Inches(op['w']),Inches(op['h']))
                tf=tb.text_frame;tf.clear();tf.word_wrap=False
                tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                fontname='Mono' if op['mono'] else ('Bold' if op['bold'] else 'Body')
                pdf.setFillColor(HexColor('#'+op['color']));pdf.setFont(fontname,op['size'])
                for li,line in enumerate(op['lines']):
                    p=tf.paragraphs[0] if li==0 else tf.add_paragraph()
                    p.text=line
                    p.font.name='Courier New' if op['mono'] else 'Arial'
                    p.font.size=Pt(op['size']);p.font.bold=op['bold'];p.font.color.rgb=RGBColor.from_string(op['color'])
                    p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=Pt(op['size']*1.22)
                    p.alignment={'left':PP_ALIGN.LEFT,'right':PP_ALIGN.RIGHT,'center':PP_ALIGN.CENTER}[op['align']]
                    xx=op['x']*72
                    if op['align']=='center': xx+=op['w']*36-pdfmetrics.stringWidth(line,fontname,op['size'])/2
                    if op['align']=='right': xx+=op['w']*72-pdfmetrics.stringWidth(line,fontname,op['size'])
                    yy=(H-op['y'])*72-op['size']*.91-li*op['size']*1.22
                    pdf.drawString(xx,yy,line)
            elif kind=='image':
                ps.shapes.add_picture(op['path'],Inches(op['x']),Inches(op['y']),Inches(op['w']),Inches(op['h']))
                pdf.drawImage(ImageReader(op['path']),op['x']*72,(H-op['y']-op['h'])*72,op['w']*72,op['h']*72,mask='auto')
            elif kind=='line':
                # The arrowhead is a small native triangle, identical in the PDF.
                from pptx.enum.shapes import MSO_CONNECTOR
                line=ps.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(op['x1']),Inches(op['y1']),Inches(op['x2']),Inches(op['y2']))
                line.line.color.rgb=RGBColor.from_string(op['color']);line.line.width=Pt(op['width'])
                pdf.setStrokeColor(HexColor('#'+op['color']));pdf.setLineWidth(op['width'])
                pdf.line(op['x1']*72,(H-op['y1'])*72,op['x2']*72,(H-op['y2'])*72)
                if op['arrow']:
                    theta=math.atan2(op['y2']-op['y1'],op['x2']-op['x1'])
                    ah=.13
                    x,y=op['x2'],op['y2']
                    tri=ps.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,Inches(x-ah*.65),Inches(y-ah*.65),Inches(ah*1.3),Inches(ah*1.3))
                    tri.rotation=math.degrees(theta)+90
                    tri.fill.solid();tri.fill.fore_color.rgb=RGBColor.from_string(op['color']);tri.line.fill.background()
                    p=pdf.beginPath();p.moveTo(x*72,(H-y)*72)
                    for a in [theta+2.65,theta-2.65]:
                        p.lineTo((x+ah*math.cos(a))*72,(H-y-ah*math.sin(a))*72)
                    p.close();pdf.setFillColor(HexColor('#'+op['color']));pdf.drawPath(p,fill=1,stroke=0)
        ps.notes_slide.notes_text_frame.text=(f'Slide {idx}: {s.title}\nPresenter {s.presenter or "Backup"}'
            +(f' | Target: {s.minutes:g} minutes' if s.minutes else '')+'\n\n'+notes.get(idx,'See speaker_notes.md.'))
        pdf.showPage()
    prs.save(ROOT/'Hastings_1970_presentation.pptx')
    pdf.save()
    serial=[dict(number=i,title=s.title,section=s.section,presenter=s.presenter,minutes=s.minutes,
                 source=s.source,backup=s.backup,content=[o.get('text',o.get('tex','')) for o in s.ops if o['kind'] in ('text','image')])
            for i,s in enumerate(slides,1)]
    (ROOT/'source'/'slide_manifest.json').write_text(json.dumps(serial,indent=2,ensure_ascii=False)+'\n')
    for i,s in enumerate(slides,1):
        for op in s.ops:
            if op['kind'] in ('text','image','rect','circle'):
                assert op['x']>=0 and op['y']>=0 and op['x']+op['w']<=W+.01 and op['y']+op['h']<=H+.01, (i,op)
    print(f'Built {len(slides)} slides; main talk {sum(s.minutes for s in slides):g} min; notes on {len(notes)} slides.')


def save_notes_pdf():
    src=ROOT/'speaker_notes.md'
    if not src.exists(): return
    style=getSampleStyleSheet()
    style.add(ParagraphStyle(name='BookTitle',fontName='Bold',fontSize=23,leading=28,textColor=HexColor('#'+NAVY),spaceAfter=18))
    style.add(ParagraphStyle(name='SlideTitle',fontName='Bold',fontSize=17,leading=21,textColor=HexColor('#'+TEAL),spaceAfter=13))
    style.add(ParagraphStyle(name='BodyEN',fontName='Body',fontSize=11,leading=16,spaceAfter=10))
    style.add(ParagraphStyle(name='MetaEN',fontName='Body',fontSize=10,leading=14,textColor=HexColor('#'+MUTED),spaceAfter=9))
    def format_inline(t):
        t=escape(t)
        t=re.sub(r'\*\*(.*?)\*\*',r'<b>\1</b>',t)
        t=re.sub(r'(?<!\*)\*([^*]+)\*',r'<i>\1</i>',t)
        t=re.sub(r'`([^`]+)`',r'<font name="Mono">\1</font>',t)
        return t.replace('\n','<br/>')
    story=[]
    for block in re.split(r'\n\s*\n',src.read_text().strip()):
        if block.startswith('# '): story.append(Paragraph(format_inline(block[2:]),style['BookTitle']))
        elif block.startswith('## '):
            if story: story.append(PageBreak())
            story.append(Paragraph(format_inline(block[3:]),style['SlideTitle']))
        elif block.startswith('### '): story.append(Paragraph(format_inline(block[4:]),style['SlideTitle']))
        elif block.startswith('---'): story.append(Spacer(1,8))
        else: story.append(Paragraph(format_inline(block),style['BodyEN']))
    def page(c,d):
        c.setStrokeColor(HexColor('#'+LINE));c.line(46,40,566,40)
        c.setFont('Body',9);c.setFillColor(HexColor('#'+MUTED))
        c.drawString(46,26,'Hastings (1970) | English speaker notes | Presenter A + Presenter B')
        c.drawRightString(566,26,str(d.page))
    doc=SimpleDocTemplate(str(ROOT/'Hastings_1970_speaker_notes.pdf'),pagesize=(612,792),
                          rightMargin=46,leftMargin=46,topMargin=45,bottomMargin=55,
                          title='Hastings (1970): English speaker notes')
    doc.build(story,onFirstPage=page,onLaterPages=page)
    print('Built printable English speaker notes.')


if __name__=='__main__':
    save_slides()
    save_notes_pdf()
